#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
W20 — 실제 이카운트 엑셀(품목마스터 + 단가마스터)에서 품목 기준정보 초안을 추출·정제한다.

원칙(오더모아 W20):
- 원본 엑셀은 저장소 밖에 두고 경로를 인자로 받는다(데이터/경로 하드코딩 금지).
- 화이트리스트 컬럼만 읽는다: 품목코드/품목명/품목구분/규격(정보)/단위/출고단가/입고단가.
  거래처/공급처/전화/주소/사업자번호 등 식별정보 컬럼은 읽지도 내보내지도 않는다.
- Git 산출물의 단가는 대표값(100원 반올림)만. 원본 exact 단가/min/max는 저장하지 않는다.
- 괄호가 공급사/브랜드로 보이면 제거, 중량·규격·형태 괄호는 유지하고 needsReview로 표시.
- 카테고리 6종(농산물/공산품/냉식/육류/수산/기타)은 키워드 초안 분류, 애매하면 기타+needsReview.
- import(DB 반영)는 하지 않는다 — 검수 가능한 초안 JSON까지만.

사용법:
  py -3 scripts/extract-real-catalog.py --items <품목마스터.xlsx> --prices <단가마스터.xlsx> --out docs/order-moa-catalog-real-draft.json
  py -3 scripts/extract-real-catalog.py --selftest   # 정제 로직 self-check
"""
from __future__ import annotations
import argparse
import json
import re
import sys
from collections import Counter
from decimal import Decimal, ROUND_HALF_UP

# 카테고리 6종 — web/src/lib/product-category.ts 와 문자열 일치(단일 소스 동기화).
CATEGORIES = ("농산물", "공산품", "냉식", "육류", "수산", "기타")
DEFAULT_CATEGORY = "기타"

# 키워드 → 카테고리(우선순위 순서대로 첫 매칭 채택). 식자재 납품 도메인 휴리스틱(초안).
CATEGORY_KEYWORDS: list[tuple[str, tuple[str, ...]]] = [
    ("공산품", ("비닐", "장갑", "락스", "세제", "봉투", "용기", "젓가락", "숟가락", "종이컵", "종이",
                "호일", "랩", "위생", "수세미", "행주", "일회용", "도시락", "트레이", "빨대", "이쑤시개",
                "앞치마", "쓰레기", "크린백", "지퍼백", "세척액", "세정", "주방", "노끈", "고무줄",
                "키친타월", "물티슈", "면장갑", "위생장갑", "받침", "꼬치",
                "롤팩", "위생백", "비닐백", "크린랩", "고무장갑", "위생모", "마스크", "채반", "소쿠리",
                "이쑤시개", "나무젓가락", "종이호일", "은박", "접시", "그릇")),
    ("육류", ("돼지", "소고기", "쇠고기", "닭", "오리", "삼겹", "목살", "등심", "안심",
              "갈비", "사태", "양지", "차돌", "항정", "가브리", "대패", "불고기", "다짐육", "곱창", "막창",
              "베이컨", "햄", "소시지", "육우", "한우", "훈제", "우삼겹", "우족", "돈까스", "돈가스",
              "우(", "돈(", "닭(", "오리(", "계육", "돈육", "우육",
              "민찌", "등뼈", "선지", "사골", "대창", "돈족", "도가니", "잡뼈", "소뼈", "돈등",
              "토시살", "부채살", "살치", "우둔", "설도", "앞다리", "뒷다리", "닭발", "똥집", "근위", "모래집",
              "소세지")),
    ("수산", ("오징어", "새우", "낙지", "문어", "주꾸미", "쭈꾸미", "조개", "홍합", "바지락", "굴비",
              "전복", "꽃게", "대게", "생선", "고등어", "갈치", "명태", "동태", "코다리", "북어", "멸치",
              "미역", "다시마", "어묵", "맛살", "게맛", "해물", "연어", "광어", "참치",
              "명란", "날치알", "오뎅", "가리비", "골뱅이", "번데기", "쥐포", "쭈",
              "젓", "꽁치", "조기", "가자미", "재첩", "생굴", "굴(", "홍어", "아귀", "병어", "서대",
              "우럭", "도미", "방어", "삼치", "임연수", "쥐치", "대구", "미더덕", "해삼", "멍게",
              "김밥김", "돌김", "돌구이김", "파래", "톳", "매생이", "황태", "적어", "장어", "메기")),
    ("냉식", ("두부", "묵", "곤약", "김치", "떡", "만두", "면", "국수", "우동", "냉면", "치즈", "계란",
              "단무지", "쫄면", "당면", "수제비", "칼국수", "유부", "어알", "동그랑", "너비아니", "떡국",
              "잡채", "묵은지", "겉절이", "깍두기", "총각")),
    ("농산물", ("나물", "채소", "버섯", "콩나물", "숙주", "오이", "상추", "배추", "대파", "쪽파",
                "양파", "마늘", "감자", "고구마", "당근", "애호박", "호박", "시금치", "부추", "깻잎",
                "고추", "미나리", "청경채", "브로콜리", "양배추", "셀러리", "생강", "도라지", "연근",
                "우엉", "토란", "더덕", "팽이", "새송이", "느타리", "표고", "양송이", "과일",
                "사과", "딸기", "토마토", "파프리카", "피망", "가지", "열무", "얼갈이", "근대",
                "아욱", "취나물", "고사리", "곤드레", "쑥갓", "청양", "홍고추", "풋고추",
                "무말랭이", "무생채", "깐마늘", "깐양파", "치커리", "케일", "루꼴라", "비트",
                "수박", "레몬", "오렌지", "포도", "파인", "멜론", "참외", "바나나", "키위", "자몽",
                "체리", "블루베리", "망고", "복숭아", "자두", "매실", "대추", "곶감", "땅콩", "견과",
                "아몬드", "호두", "잣", "은행", "총각무", "알타리", "봄동", "아스파라거스", "깐밤",
                "모듬씨앗", "옥수수", "로메인", "샐러리", "무순", "새싹", "숙주나물", "얼갈이배추")),
]

# 괄호 안에서 '유지할' 규격/원산지/형태 토큰. 이 목록에 없으면 브랜드/공급사로 보고 제거한다(화이트리스트).
FORM_WORDS = (
    "국산", "국내산", "국내", "수입", "중국산", "국내재배", "자연산", "양식", "원양",
    "냉장", "냉동", "생물", "건조", "자숙", "생", "냉", "진공", "세척", "손질",
    "특", "왕특", "상", "중", "소", "대", "특상", "상품",
    "채", "찌개", "전골", "국거리", "불고기", "목전지", "까스", "가스", "등심", "안심",
    "사태", "양지", "갈비", "다짐", "다진", "반달", "염장", "절임", "국물", "육수",
    "볶음", "무침", "구이", "조림", "낱개", "낱", "묶음", "실", "포찹", "절단",
    "슬라이스", "분태", "파지", "주황", "색", "알", "순", "즙", "반찬", "손잡이", "완",
)
UNIT_TOKEN = re.compile(r"\b(\d+\s*)?(kg|kilo|g|box|봉지|봉|판|망|개입|개|단|팩|통|장|모|ea|리터|리|ml|포기|포|근|되|짝|줄|입|시루|박스|병|캔|롤|말|자루|묶음)\b", re.I)
UNIT_NORM = {"kilo": "kg", "박스": "BOX", "box": "BOX", "봉지": "봉", "개입": "개"}

def is_spec_token(t: str) -> bool:
    """괄호 안 한 토큰이 규격/원산지/형태(유지 대상)인가. 아니면 브랜드/공급사로 간주(제거)."""
    t = t.strip()
    if not t:
        return False
    if re.search(r"\d", t) and re.fullmatch(r"[\d.]+\s*[a-zA-Z가-힣]{0,5}", t):
        return True  # 1kg, 100L, 50ea, 2kg, 10미, 100매 …
    if re.fullmatch(r"[a-zA-Z]{1,5}", t):
        return True  # BOX, EA, L …
    return any(w in t for w in FORM_WORDS)

# 식별정보 자체 스캔용 — 공급사 접미사/전화/사업자번호만(품목 설명어 정육·반찬 등은 제외해 오탐 최소화).
IDENTITY_RE = re.compile(r"(식품|상회|유통|축산|수산|농산|물산|상사|유업|농장|산업|F&B|[가-힣]{2,}상회|0\d{1,2}[-\s]?\d{3,4}[-\s]?\d{4}|\d{3}-\d{2}-\d{5})")

# 키워드 단순 포함으로 잘못 분류되기 쉬운 복합어. 검수 대상 기타로 남긴다.
AMBIGUOUS_CATEGORY_TERMS = ("분모자", "고추장", "쌈장", "된장", "곱창김")


def round_price(v) -> int | None:
    """대표단가 = 100원 단위 반올림. 없음/0/음수 → None(needsReview)."""
    try:
        n = float(v)
    except (TypeError, ValueError):
        return None
    if n <= 0:
        return None
    return int((Decimal(str(n)) / Decimal("100")).quantize(Decimal("1"), rounding=ROUND_HALF_UP) * 100)


def clean_name(raw: str) -> tuple[str, bool, str | None]:
    """품목명 정제. 반환 (정제명, needsReview, 사유|None).
    괄호 안을 [,/]로 쪼개 규격/원산지/형태 토큰만 유지하고, 그 외(브랜드·공급사·미상)는 제거한다.
    무언가 제거되면 needsReview(괄호정보정리) — 식별정보가 이름에 남지 않게 하는 화이트리스트 방식.
    """
    name = re.sub(r"\s+", " ", str(raw)).strip()
    dropped: list[str] = []
    def repl(m):
        parts = [p.strip() for p in re.split(r"[,/]", m.group(1)) if p.strip()]
        kept = [p for p in parts if is_spec_token(p)]
        dropped.extend(p for p in parts if not is_spec_token(p))
        return f"({','.join(kept)})" if kept else ""
    cleaned = re.sub(r"\(([^)]*)\)", repl, name).strip()
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    if dropped:
        return (cleaned or name), True, "괄호정보정리"
    return (cleaned or name), False, None


def derive_unit(unit_col, spec: str) -> tuple[str | None, bool]:
    """단위 결정: 단위 컬럼 우선, 없으면 규격에서 단위 토큰 추출. 반환 (단위|None, needsReview)."""
    if unit_col and str(unit_col).strip():
        u = str(unit_col).strip()
        return UNIT_NORM.get(u.lower(), u), False
    if spec:
        m = UNIT_TOKEN.search(str(spec))
        if m:
            u = m.group(2)
            return UNIT_NORM.get(u.lower(), u), False
    return None, True


def classify_category(name: str, spec: str) -> tuple[str, str | None]:
    """키워드 초안 분류. 반환 (카테고리, 매칭키워드|None). 미매칭 → (기타, None)."""
    hay = f"{name} {spec or ''}"
    if any(term in hay for term in AMBIGUOUS_CATEGORY_TERMS):
        return DEFAULT_CATEGORY, None
    for cat, kws in CATEGORY_KEYWORDS:
        for kw in kws:
            if kw in hay:
                return cat, kw
    return DEFAULT_CATEGORY, None


def alias_candidates(cleaned: str) -> list[str]:
    """별칭 후보: 괄호/대괄호 앞 짧은 핵심어 등(검색창내용이 비어 있어 이름 기반)."""
    base = re.sub(r"[\(\[].*", "", cleaned).strip()
    cands = set()
    if base and base != cleaned:
        cands.add(base)
    return sorted(c for c in cands if 1 <= len(c) <= 12 and c != cleaned)


# ---------- 엑셀 읽기(화이트리스트) ----------

def _load_sheet(path: str, want_cols: set[str]):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb.worksheets[0]
    # 헤더 행 = 품목코드 & 품목명 포함 행
    hdr_row, headers = None, []
    for i in range(1, 8):
        vals = [ws.cell(row=i, column=j).value for j in range(1, (ws.max_column or 1) + 1)]
        strs = [str(v).strip() if v is not None else "" for v in vals]
        if "품목코드" in strs and "품목명" in strs:
            hdr_row, headers = i, strs
            break
    if hdr_row is None:
        raise SystemExit(f"[오류] {path}: 헤더 행(품목코드·품목명)을 못 찾음")
    idx = {h: k for k, h in enumerate(headers)}
    out = []
    for row in ws.iter_rows(min_row=hdr_row + 1, max_row=ws.max_row, values_only=True):
        rec = {}
        for c in want_cols:
            if c in idx and idx[c] < len(row):
                rec[c] = row[idx[c]]
        if rec.get("품목코드") is not None:
            out.append(rec)
    wb.close()
    return out, headers


def extract(items_path: str, prices_path: str):
    items, item_headers = _load_sheet(items_path, {"품목코드", "품목명", "품목구분", "규격정보"})
    prices, price_headers = _load_sheet(prices_path, {"품목코드", "품목명", "규격", "단위", "출고단가", "입고단가"})
    by_code_item = {str(r["품목코드"]).strip(): r for r in items}

    excluded_columns = sorted(set(item_headers + price_headers) - {
        "품목코드", "품목명", "품목구분", "규격정보", "규격", "단위", "출고단가", "입고단가"
    })

    catalog = []
    for pr in prices:
        code = str(pr["품목코드"]).strip()
        raw_name = pr.get("품목명") or by_code_item.get(code, {}).get("품목명") or ""
        if not str(raw_name).strip():
            continue
        spec = (pr.get("규격") or by_code_item.get(code, {}).get("규격정보") or "")
        spec = re.sub(r"\s+", " ", str(spec)).strip()
        if str(raw_name).strip() in {"품목", "품목명"} and not spec and not pr.get("단위") and not pr.get("출고단가"):
            continue
        name, name_rev, name_reason = clean_name(raw_name)
        unit, unit_rev = derive_unit(pr.get("단위"), spec)
        cat, kw = classify_category(name, spec)
        out_price = round_price(pr.get("출고단가"))
        in_price = round_price(pr.get("입고단가"))

        reasons = []
        if name_rev and name_reason:
            reasons.append(name_reason)
        if unit_rev:
            reasons.append("단위없음")
        if out_price is None:
            reasons.append("출고단가없음")
        if kw is None:
            reasons.append("카테고리미분류")

        catalog.append({
            "code": code,
            "name": name,
            "spec": spec or None,
            "unit": unit,
            "category": cat,
            "categoryMatched": kw,
            "repSalePrice": out_price,      # 대표 출고단가(100원 반올림)
            "repPurchasePrice": in_price,   # 대표 입고단가(100원 반올림)
            "aliasCandidates": alias_candidates(name),
            "needsReview": bool(reasons),
            "reviewReasons": reasons,
        })
    catalog.sort(key=lambda r: r["code"])  # 멱등: 같은 입력 → 같은 출력
    return catalog, excluded_columns


def build_report(catalog, excluded_columns):
    cat_dist = Counter(c["category"] for c in catalog)
    unit_dist = Counter((c["unit"] or "단위없음") for c in catalog)
    review = [c for c in catalog if c["needsReview"]]
    reason_dist = Counter(r for c in review for r in c["reviewReasons"])
    names = [c["name"] for c in catalog]
    dup = Counter(names)
    multi = {n: k for n, k in dup.items() if k > 1}
    return {
        "totalItems": len(catalog),
        "uniqueNames": len(set(names)),
        "excludedColumns": excluded_columns,
        "categoryDist": dict(cat_dist.most_common()),
        "unitDistTop": dict(unit_dist.most_common(15)),
        "needsReviewCount": len(review),
        "needsReviewReasons": dict(reason_dist.most_common()),
        "autoClassifiedPct": round(100 * (len(catalog) - cat_dist.get("기타", 0)) / max(1, len(catalog)), 1),
        "priceMissingCount": sum(1 for c in catalog if c["repSalePrice"] is None),
        "duplicateNameGroups": len(multi),          # 같은 품목명 다규격/다단위(=Phase3 다단위 신호)
        "duplicateNameItems": sum(multi.values()),
        "aliasCandidateCount": sum(1 for c in catalog if c["aliasCandidates"]),
    }


def identity_scan(catalog) -> list[str]:
    """산출물에 식별정보 의심값(상호/전화/사업자번호)이 새지 않았는지 자체 스캔."""
    hits = []
    for c in catalog:
        blob = " ".join(str(c.get(k) or "") for k in ("name", "spec"))
        for m in IDENTITY_RE.findall(blob):
            hits.append(f"{c['code']}::{c['name']}::{m}")
    return hits


def run_selftest():
    # round_price
    assert round_price(1000) == 1000
    assert round_price(13040) == 13000 and round_price(13050) == 13100
    assert round_price(0) is None and round_price(None) is None and round_price("") is None
    # clean_name: 브랜드/공급사 토큰 제거(규격/원산지/형태는 유지). 식별정보 유출 방지 핵심.
    assert clean_name("세척숙주(성진식품)")[:2] == ("세척숙주", True)
    assert clean_name("다진오징어(합천수산/1kg)")[:2] == ("다진오징어(1kg)", True)   # 공급사 제거, 규격 유지
    assert clean_name("고춧가루(장용,수입,별식품)")[:2] == ("고춧가루(수입)", True)   # 원산지만 유지
    assert clean_name("커피믹스(프렌치카페)")[:2] == ("커피믹스", True)               # 브랜드 통째 제거
    assert clean_name("돈(찌개)")[:2] == ("돈(찌개)", False)                          # 형태어는 유지·검토불필요
    assert clean_name("콩나물")[:2] == ("콩나물", False)
    # derive_unit: 단위 컬럼 우선, 없으면 규격에서
    assert derive_unit("BOX", "10KG 국산") == ("BOX", False)
    assert derive_unit(None, "1kg 국산")[0] == "kg"
    assert derive_unit(None, "국내산") == (None, True)
    # classify_category
    assert classify_category("콩나물", "1kg")[0] == "농산물"
    assert classify_category("돈(찌개)", "1kg")[0] == "육류"
    assert classify_category("다진오징어", "1kg")[0] == "수산"
    assert classify_category("위생장갑", "100매")[0] == "공산품"
    assert classify_category("팩두부", "3kg")[0] == "냉식"
    assert classify_category("곱창김", "100장")[0] == "기타"
    assert classify_category("고추장", "3kg")[0] == "기타"
    assert classify_category("수박분모자", "500g*20")[0] == "기타"
    assert classify_category("비엔나소세지", "1kg")[0] == "육류"
    assert classify_category("듣도보도못한것XYZ", "")[0] == "기타"
    print("self-check OK")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--items", help="품목마스터 xlsx 경로(저장소 밖)")
    ap.add_argument("--prices", help="단가마스터 xlsx 경로(저장소 밖)")
    ap.add_argument("--out", help="산출물 JSON 경로")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        run_selftest()
        return
    if not (a.items and a.prices and a.out):
        ap.error("--items --prices --out 필요 (또는 --selftest)")
    catalog, excluded = extract(a.items, a.prices)
    report = build_report(catalog, excluded)
    hits = identity_scan(catalog)
    payload = {
        "meta": {
            "source": "실 이카운트 품목마스터+단가마스터(저장소 밖, 미커밋). 식별정보 미수록.",
            "note": "W20 품목 기준정보 초안. 단가는 100원 반올림 대표값(exact 아님). import(DB 반영)는 W21 검수 후.",
            "categories": list(CATEGORIES),
        },
        "report": report,
        "identitySuspects": hits,   # 비어 있어야 정상
        "items": catalog,
    }
    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    print(f"\n식별정보 의심값: {len(hits)}건")
    for h in hits[:20]:
        print("  ", h)
    print(f"\n산출물: {a.out}")


if __name__ == "__main__":
    main()
